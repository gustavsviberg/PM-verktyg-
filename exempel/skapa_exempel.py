import docx, openpyxl
from docx.shared import Pt
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, Spacer, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
OUT="/mnt/project-files/pm-verktyg/exempel/"
roles=[("Uppdragsledare",1250),("Konstruktör",1050),("Handläggare/ritare",900),("BIM-samordnare",1000)]
rp=dict(roles)
phases=[("Fas 1 – Utredning",[("1.1","Startmöte och inventering av befintligt underlag","Uppdragsledare",12),("1.2","Förstudie stomval","Konstruktör",40)]),
("Fas 2 – Systemhandling",[("2.1","Systemhandling konstruktion","Konstruktör",120),("2.2","Projekteringsmöten (12 st)","Uppdragsledare",36)]),
("Fas 3 – Bygghandling",[("3.1","Bygghandling stomme och grund","Konstruktör",220),("3.2","Ritningsproduktion","Handläggare/ritare",180),("3.3","BIM-samordning och kollisionskontroll","BIM-samordnare",70),("3.4","Egenkontroll och granskning","Uppdragsledare",16)]),
("Fas 4 – Byggskede",[("4.1","Medverkan i byggskede, byggplatsbesök","Konstruktör",48),("4.2","Relationshandlingar","Handläggare/ritare",30)])]
def fmt(n): return f"{n:,.0f}".replace(","," ")
total=sum(h*rp[r] for _,ts in phases for _,_,r,h in ts); print("total",total)
intro=["Anbud A-2026-114","Anbudsdatum: 2026-10-20",
"Till: Lindvik kommun, Fastighetsavdelningen","Avser: Konstruktionsprojektering, nybyggnad Ekbackens förskola",
"Vi tackar för förfrågan och lämnar härmed anbud på konstruktionsprojektering för nybyggnad av Ekbackens förskola, två plan med stomme i KL-trä och grundläggning på platta på mark.",
"Uppdraget planeras att starta 2026-11-16 och avslutas 2027-08-31 i och med överlämnade relationshandlingar."]
arvode=[f"Uppdraget utförs till fast arvode om {fmt(total)} kr exklusive moms, enligt ABK 09.",
f"Reserv för oförutsett: 5 % ({fmt(total*0.05)} kr). Reserven ingår inte i det fasta arvodet och avropas efter beställarens skriftliga godkännande.",
"Á-priser: Tillkommande projekteringsmöte utöver 12 st debiteras med 4 500 kr/st. Tilläggsarbeten (ÄTA) debiteras enligt timpriserna ovan.",
"Fakturering sker månadsvis i takt med upparbetning. Betalningsvillkor 30 dagar.","Med vänlig hälsning","Karin Ström, Uppdragsledare, Stomkonsult Norr AB"]
# DOCX
d=docx.Document(); d.add_heading("Anbud – Konstruktionsprojektering, nybyggnad Ekbackens förskola",1)
for p in intro: d.add_paragraph(p)
d.add_heading("Timpriser",2); t=d.add_table(rows=1,cols=2); t.style="Table Grid"
t.rows[0].cells[0].text="Roll"; t.rows[0].cells[1].text="Timpris (kr/h)"
for r,p in roles: c=t.add_row().cells; c[0].text=r; c[1].text=fmt(p)
d.add_heading("Arbetsmoment och leveranser",2); t=d.add_table(rows=1,cols=5); t.style="Table Grid"
for i,h in enumerate(["Kod","Moment","Roll","Timmar","Belopp (kr)"]): t.rows[0].cells[i].text=h
for ph,ts in phases:
    c=t.add_row().cells; c[0].text=ph
    for k,n,r,h in ts:
        c=t.add_row().cells; c[0].text=k; c[1].text=n; c[2].text=r; c[3].text=str(h); c[4].text=fmt(h*rp[r])
c=t.add_row().cells; c[1].text="Summa"; c[3].text=str(sum(h for _,ts in phases for *_,h in ts)); c[4].text=fmt(total)
d.add_heading("Arvode och villkor",2)
for p in arvode: d.add_paragraph(p)
d.save(OUT+"Exempelanbud_Ekbackens_forskola.docx")
# PDF
ss=getSampleStyleSheet(); el=[Paragraph("Anbud – Konstruktionsprojektering, nybyggnad Ekbackens förskola",ss["Title"])]
el+= [Paragraph(p,ss["Normal"]) for p in intro]; el.append(Spacer(1,10)); el.append(Paragraph("Timpriser",ss["Heading2"]))
ts_=TableStyle([("GRID",(0,0),(-1,-1),0.4,colors.grey),("FONTSIZE",(0,0),(-1,-1),8.5)])
el.append(Table([["Roll","Timpris (kr/h)"]]+[[r,fmt(p)] for r,p in roles],style=ts_))
el.append(Paragraph("Arbetsmoment och leveranser",ss["Heading2"])); rows=[["Kod","Moment","Roll","Timmar","Belopp (kr)"]]
for ph,ts in phases:
    rows.append([ph,"","","",""]); rows+= [[k,n,r,str(h),fmt(h*rp[r])] for k,n,r,h in ts]
rows.append(["","Summa","",str(sum(h for _,ts in phases for *_,h in ts)),fmt(total)])
el.append(Table(rows,style=ts_)); el.append(Paragraph("Arvode och villkor",ss["Heading2"])); el+=[Paragraph(p,ss["Normal"]) for p in arvode]
SimpleDocTemplate(OUT+"Exempelanbud_Ekbackens_forskola.pdf",pagesize=A4).build(el)
# XLSX ambiguous
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Anbud"
B=[("Anbud","Dagvattenutredning Norra hamnen"),("Beställare","Säterö Hamn AB"),("Anbudsnummer","2026-0871"),("Anbudsdatum","2026-10-12"),("Period","2026-12-01 – 2027-04-30"),("",""),("Roll","Timpris (kr/h)")]
for r in B: ws.append(list(r))
r2=[("Uppdragsledare",1300),("Utredare",1080),("GIS-ingenjör",980)]; rp2=dict(r2)
for r in r2: ws.append(list(r))
ws.append([]); ws.append(["Nr","Moment","Roll","Timmar","Belopp (kr)"])
t2=[("1","Insamling av underlag","Utredare",24),("2","Platsbesök och inmätning","GIS-ingenjör",32),("3","Hydraulisk modellering","Utredare",90),("4","Åtgärdsförslag","Utredare",40),("5","Rapport och presentation","Uppdragsledare",28),("6","Möten med beställaren","Uppdragsledare",12)]
for k,n,r,h in t2: ws.append([k,n,r,h,h*rp2[r]])
tot2=sum(h*rp2[r] for *_,r,h in t2); ws.append(["","Bedömd kostnad","",sum(x[3] for x in t2),tot2]); print("tot2",tot2)
ws.append([]); ws.append(["Villkor","Ersättning enligt ABK 09. Utskrift av rapport á 650 kr/st."])
for c in ("A","B","C"): ws.column_dimensions[c].width=28
for row in ws.iter_rows():
    if row[0].value in ("Roll","Nr"):
        for c in row: c.font=Font(bold=True)
wb.save(OUT+"Exempelanbud_Dagvatten_Norra_hamnen.xlsx")
